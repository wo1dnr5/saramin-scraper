import csv
import re
from collections import Counter, defaultdict
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

INPUT  = 'output/saramin_jobs_data.csv'
OUTPUT = 'output/채용공고_서울.xlsx'

# ── 색상 ─────────────────────────────────────────────────
C_HEADER_BG  = '1F3864'
C_HEADER_FG  = 'FFFFFF'
C_EVEN_BG    = 'EBF3FB'
C_TITLE_BG   = '2E75B6'
C_BORDER     = 'BDD7EE'
C_LINK       = '1155CC'

def side(color=C_BORDER):
    s = Side(style='thin', color=color)
    return Border(left=s, right=s, top=s, bottom=s)

def fill(hex_color):
    return PatternFill('solid', fgColor=hex_color)

with open(INPUT, encoding='utf-8-sig') as f:
    rows = list(csv.DictReader(f))

wb = Workbook()

# ══════════════════════════════════════════════════════════
# 시트 1 — 전체 공고 목록
# ══════════════════════════════════════════════════════════
ws = wb.active
ws.title = '📋 전체 공고'

# 제목 행
ws.merge_cells('A1:L1')
tc = ws['A1']
tc.value = '서울 채용공고  |  솔루션엔지니어 · 데이터분석가 · 시스템엔지니어  |  사람인 수집'
tc.font = Font(name='맑은 고딕', bold=True, size=13, color='FFFFFF')
tc.fill = fill(C_TITLE_BG)
tc.alignment = Alignment(horizontal='center', vertical='center')
ws.row_dimensions[1].height = 30

# 요약 행
ws.merge_cells('A2:L2')
sc = ws['A2']
sc.value = f'총 {len(rows)}개 공고   |   수집일: 2026-04-22   |   지역: 서울'
sc.font = Font(name='맑은 고딕', size=10, color='595959')
sc.fill = fill('EBF3FB')
sc.alignment = Alignment(horizontal='left', vertical='center', indent=1)
ws.row_dimensions[2].height = 18

# 헤더
COLS = [
    ('No.',        4),
    ('검색어',     12),
    ('기업명',     22),
    ('공고제목',   40),
    ('고용형태',   10),
    ('경력',       13),
    ('학력',       10),
    ('지역',       14),
    ('마감일',     12),
    ('직무카테고리', 30),
    ('특이사항',   18),
    ('공고링크',   10),
]

for ci, (h, w) in enumerate(COLS, 1):
    c = ws.cell(row=3, column=ci, value=h)
    c.font      = Font(name='맑은 고딕', bold=True, size=11, color=C_HEADER_FG)
    c.fill      = fill(C_HEADER_BG)
    c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    c.border    = side('FFFFFF')
    ws.column_dimensions[get_column_letter(ci)].width = w
ws.row_dimensions[3].height = 26

# 데이터
for i, row in enumerate(rows):
    r = i + 4
    bg = 'FFFFFF' if i % 2 == 0 else C_EVEN_BG

    job_type = row.get('고용형태', '') or ''
    badge    = row.get('배지', '') or ''

    values = [
        i + 1,
        row.get('검색어', ''),
        row.get('기업명', ''),
        row.get('공고제목', ''),
        job_type,
        row.get('경력', ''),
        row.get('학력', ''),
        row.get('지역', ''),
        row.get('마감일', ''),
        row.get('직무카테고리', ''),
        badge,
        '바로가기',
    ]

    for ci, val in enumerate(values, 1):
        c = ws.cell(row=r, column=ci, value=val)
        c.font      = Font(name='맑은 고딕', size=10)
        c.fill      = fill(bg)
        c.border    = side()
        c.alignment = Alignment(vertical='center', wrap_text=True)
        if ci in (1, 5, 7, 9, 12):
            c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

    # 고용형태 색상
    jt_cell = ws.cell(row=r, column=5)
    if '정규직' in job_type:
        jt_cell.fill = fill('D9EAD3')
        jt_cell.font = Font(name='맑은 고딕', size=10, bold=True, color='274E13')
    elif '계약직' in job_type or '기간제' in job_type:
        jt_cell.fill = fill('FCE5CD')
        jt_cell.font = Font(name='맑은 고딕', size=10, bold=True, color='7F2B00')

    # 특이사항(배지) 색상
    if badge:
        ws.cell(row=r, column=11).font = Font(name='맑은 고딕', size=10, bold=True, color='7030A0')

    # 링크
    link = row.get('공고링크', '')
    if link and link != 'N/A':
        lc = ws.cell(row=r, column=12)
        lc.hyperlink = link
        lc.value = '바로가기 →'
        lc.font = Font(name='맑은 고딕', size=10, color=C_LINK, underline='single')

    ws.row_dimensions[r].height = 40

ws.auto_filter.ref = f'A3:L{3 + len(rows)}'
ws.freeze_panes = 'A4'


# ══════════════════════════════════════════════════════════
# 시트 2 — 검색어별 공고 목록 (3개 키워드별 탭)
# ══════════════════════════════════════════════════════════
keyword_map = defaultdict(list)
for row in rows:
    keyword_map[row.get('검색어', '기타')].append(row)

KEYWORD_COLORS = {
    '솔루션엔지니어': ('1F3864', 'D6E4F7'),
    '데이터분석가':   ('375623', 'E2EFDA'),
    '시스템엔지니어': ('3F3151', 'EAD9F7'),
}

for kw, kw_rows in keyword_map.items():
    header_c, even_c = KEYWORD_COLORS.get(kw, ('333333', 'F2F2F2'))
    wsk = wb.create_sheet(f'🔍 {kw}')

    wsk.merge_cells('A1:J1')
    tk = wsk['A1']
    tk.value = f'{kw}  |  총 {len(kw_rows)}개 공고  |  서울'
    tk.font = Font(name='맑은 고딕', bold=True, size=13, color='FFFFFF')
    tk.fill = fill(header_c)
    tk.alignment = Alignment(horizontal='center', vertical='center')
    wsk.row_dimensions[1].height = 28

    KW_COLS = [
        ('No.',      4), ('기업명', 22), ('공고제목', 42), ('고용형태', 10),
        ('경력', 13), ('학력', 10), ('지역', 14), ('마감일', 12),
        ('특이사항', 18), ('공고링크', 10),
    ]
    for ci, (h, w) in enumerate(KW_COLS, 1):
        c = wsk.cell(row=2, column=ci, value=h)
        c.font      = Font(name='맑은 고딕', bold=True, size=11, color='FFFFFF')
        c.fill      = fill(header_c)
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        c.border    = side('FFFFFF')
        wsk.column_dimensions[get_column_letter(ci)].width = w
    wsk.row_dimensions[2].height = 24

    for i, row in enumerate(kw_rows):
        r = i + 3
        bg = 'FFFFFF' if i % 2 == 0 else even_c
        job_type = row.get('고용형태', '') or ''
        badge    = row.get('배지', '') or ''
        vals = [
            i + 1,
            row.get('기업명', ''),
            row.get('공고제목', ''),
            job_type,
            row.get('경력', ''),
            row.get('학력', ''),
            row.get('지역', ''),
            row.get('마감일', ''),
            badge,
            '바로가기',
        ]
        for ci, val in enumerate(vals, 1):
            c = wsk.cell(row=r, column=ci, value=val)
            c.font      = Font(name='맑은 고딕', size=10)
            c.fill      = fill(bg)
            c.border    = side()
            c.alignment = Alignment(vertical='center', wrap_text=True)
            if ci in (1, 4, 6, 8, 10):
                c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

        jt_cell = wsk.cell(row=r, column=4)
        if '정규직' in job_type:
            jt_cell.fill = fill('D9EAD3')
            jt_cell.font = Font(name='맑은 고딕', size=10, bold=True, color='274E13')
        elif '계약직' in job_type or '기간제' in job_type:
            jt_cell.fill = fill('FCE5CD')
            jt_cell.font = Font(name='맑은 고딕', size=10, bold=True, color='7F2B00')

        link = row.get('공고링크', '')
        if link and link != 'N/A':
            lc = wsk.cell(row=r, column=10)
            lc.hyperlink = link
            lc.value = '바로가기 →'
            lc.font = Font(name='맑은 고딕', size=10, color=C_LINK, underline='single')

        wsk.row_dimensions[r].height = 40

    wsk.auto_filter.ref = f'A2:J{2 + len(kw_rows)}'
    wsk.freeze_panes = 'A3'


# ══════════════════════════════════════════════════════════
# 시트 3 — 통계
# ══════════════════════════════════════════════════════════
ws3 = wb.create_sheet('📊 통계')

ws3.merge_cells('A1:F1')
t3 = ws3['A1']
t3.value = '채용공고 통계 요약'
t3.font = Font(name='맑은 고딕', bold=True, size=13, color='FFFFFF')
t3.fill = fill('1F3864')
t3.alignment = Alignment(horizontal='center', vertical='center')
ws3.row_dimensions[1].height = 28

def stat_block(ws, start_row, col, title, data):
    ws.merge_cells(start_row=start_row, end_row=start_row, start_column=col, end_column=col+2)
    tc = ws.cell(row=start_row, column=col, value=title)
    tc.font = Font(name='맑은 고딕', bold=True, size=11, color='FFFFFF')
    tc.fill = fill('2E75B6')
    tc.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[start_row].height = 22

    for ci, h in enumerate(['항목', '건수', '비율(%)'], col):
        c = ws.cell(row=start_row+1, column=ci, value=h)
        c.font = Font(name='맑은 고딕', bold=True, size=10)
        c.fill = fill('D6E4F7')
        c.alignment = Alignment(horizontal='center', vertical='center')
        c.border = side()
    ws.row_dimensions[start_row+1].height = 20

    total = sum(v for _, v in data)
    for i, (label, cnt) in enumerate(data):
        ri = start_row + 2 + i
        pct = round(cnt / total * 100, 1) if total else 0
        for ci, v in zip(range(col, col+3), [label, cnt, pct]):
            c = ws.cell(row=ri, column=ci, value=v)
            c.font = Font(name='맑은 고딕', size=10)
            c.fill = fill('F2F7FD' if i % 2 == 0 else 'FFFFFF')
            c.border = side()
            c.alignment = Alignment(
                horizontal='center' if ci > col else 'left',
                vertical='center'
            )
        ws.row_dimensions[ri].height = 20

# 검색어별
kw_cnt = Counter(r.get('검색어', '') for r in rows)
stat_block(ws3, 3, 1, '검색어별 공고 수', kw_cnt.most_common())

# 고용형태별
jt_cnt = Counter(r.get('고용형태', '미기재') or '미기재' for r in rows)
stat_block(ws3, 3, 5, '고용형태별 공고 수', jt_cnt.most_common())

# 경력별
exp_cnt = Counter(r.get('경력', '미기재') or '미기재' for r in rows)
stat_block(ws3, 10, 1, '경력별 공고 수', exp_cnt.most_common())

# 지역별 (자치구)
gu_cnt = Counter()
for row in rows:
    loc = re.sub(r'^서울\s*', '', row.get('지역', '') or '')
    gu_cnt[loc if loc else '서울전체'] += 1
stat_block(ws3, 10, 5, '지역별 공고 수', gu_cnt.most_common())

for ci, w in enumerate([18, 8, 8, 2, 18, 8, 8], 1):
    ws3.column_dimensions[get_column_letter(ci)].width = w
ws3.freeze_panes = 'A2'

# ── 저장 ─────────────────────────────────────────────────
wb.save(OUTPUT)
print(f'저장 완료 → {OUTPUT}')
print(f'총 {len(rows)}개 공고')
